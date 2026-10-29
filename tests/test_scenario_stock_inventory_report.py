import datetime as dt
import unittest
from unittest.mock import patch
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pypdf import PdfReader
from decimal import Decimal

from proteus import Model, Wizard
from trytond.modules.company.tests.tools import create_company, get_company
from trytond.modules.currency.tests.tools import get_currency
from trytond.modules.product import price_digits
from trytond.modules.stock_inventory_jreport import inventory as report_module
from trytond.tests.test_tryton import drop_db
from trytond.tests.tools import activate_modules


class TestStockInventoryReport(unittest.TestCase):

    def setUp(self):
        drop_db()
        super().setUp()

    def tearDown(self):
        drop_db()
        super().tearDown()

    def test(self):
        activate_modules(['stock_inventory_jreport', 'stock_valued',
            'stock_lot_cost'])

        create_company()
        get_company()
        currency = get_currency()

        ProductUom = Model.get('product.uom')
        unit, = ProductUom.find([('name', '=', 'Unit')])

        ProductTemplate = Model.get('product.template')

        template1 = ProductTemplate()
        template1.name = 'Product 1'
        template1.code = 'P1'
        template1.default_uom = unit
        template1.type = 'goods'
        template1.list_price = Decimal('20')
        template1.cost_price_method = 'fixed'
        template1.save()
        product1, = template1.products
        product1.cost_price = Decimal('5')
        product1.save()

        template2 = ProductTemplate()
        template2.name = 'Product 2'
        template2.code = 'P2'
        template2.default_uom = unit
        template2.type = 'goods'
        template2.list_price = Decimal('20')
        template2.cost_price_method = 'fixed'
        template2.save()
        product2, = template2.products
        product2.cost_price = Decimal('7')
        product2.save()

        Location = Model.get('stock.location')
        warehouse, = Location.find([('code', '=', 'WH')])
        storage_loc, = Location.find([('code', '=', 'STO')])
        supplier_loc, = Location.find([('code', '=', 'SUP')])

        StockMove = Model.get('stock.move')
        today = dt.date.today()

        move1 = StockMove()
        move1.product = product1
        move1.unit = product1.default_uom
        move1.quantity = 5
        move1.from_location = supplier_loc
        move1.to_location = storage_loc
        move1.planned_date = today
        move1.effective_date = today
        move1.unit_price = Decimal('5')
        move1.currency = currency

        move2 = StockMove()
        move2.product = product2
        move2.unit = product2.default_uom
        move2.quantity = 3
        move2.from_location = supplier_loc
        move2.to_location = storage_loc
        move2.planned_date = today
        move2.effective_date = today
        move2.unit_price = Decimal('7')
        move2.currency = currency
        StockMove.click([move1, move2], 'do')

        Inventory = Model.get('stock.inventory')
        inventory = Inventory()
        inventory.location = storage_loc
        inventory.empty_quantity = 'keep'
        inventory.save()
        inventory.click('complete_lines')
        lines = {line.product.id: line for line in inventory.lines}
        lines[product1.id].quantity = 4
        lines[product2.id].quantity = 3
        inventory.save()
        inventory.click('confirm')

        Product = Model.get('product.product')
        wizard = Wizard('stock.inventory.print_total_inventory')
        self.assertEqual(wizard.form.output_format, 'pdf')
        self.assertEqual(wizard.form.order, 'location')
        wizard.form.locations.append(Location(warehouse.id))
        wizard.form.products.append(Product(product1.id))
        wizard.form.products.append(Product(product2.id))
        wizard.form.quantities = 'positive'
        wizard.form.order = 'location'
        wizard.form.output_format = 'pdf'
        wizard.form.timeout = 120
        wizard.execute('print_')
        self.assertEqual(wizard.state, 'end')
        self.assertEqual(len(wizard.actions), 1)
        oext, content, _, _ = wizard.actions[0]
        self.assertEqual(oext, 'pdf')
        self.assertIsInstance(content, bytes)

        wizard = Wizard('stock.inventory.print_total_inventory')
        wizard.form.locations.append(Location(warehouse.id))
        wizard.form.products.append(Product(product1.id))
        wizard.form.products.append(Product(product2.id))
        wizard.form.quantities = 'positive'
        wizard.form.order = 'location'
        wizard.form.output_format = 'xlsx'
        wizard.form.timeout = 120
        wizard.execute('print_')
        self.assertEqual(wizard.state, 'end')
        self.assertEqual(len(wizard.actions), 1)
        oext, content, _, _ = wizard.actions[0]
        self.assertEqual(oext, 'xlsx')
        self.assertIsInstance(content, bytes)
        self.assertTrue(content.startswith(b'PK'))

        sheet = load_workbook(BytesIO(content)).active
        rows = list(sheet.values)
        header = next(row for row in rows if row[0] == 'Location')
        self.assertEqual(header, ('Location', 'Product', 'Quantity'))
        quantities = {row[1]: row[2] for row in rows
            if row[1] in (product1.rec_name, product2.rec_name)}
        self.assertEqual(quantities, {product1.rec_name: 4, product2.rec_name: 3})

        Lot = Model.get('stock.lot')
        lot = Lot(number='LOT-VALUED', product=product2)
        lot.save()
        move = StockMove(product=product2, unit=unit, quantity=2,
            from_location=supplier_loc, to_location=storage_loc,
            planned_date=today, effective_date=today,
            unit_price=Decimal('11'), currency=currency, lot=lot)
        move.click('do')

        value_format = '#,##0' + (
            '.' + '0' * price_digits[1] if price_digits[1] else '')
        for by_lot in [False, True]:
            for output_format in ['xlsx', 'pdf']:
                wizard = Wizard('stock.inventory.print_total_inventory')
                wizard.form.locations.append(Location(warehouse.id))
                wizard.form.valued = True
                wizard.form.group_by_lot = by_lot
                wizard.form.output_format = output_format
                wizard.execute('print_')
                self.assertEqual(wizard.state, 'end')
                self.assertEqual(len(wizard.actions), 1)
                extension, content, _, _ = wizard.actions[0]
                self.assertEqual(extension, output_format)
                if output_format == 'pdf':
                    self.assertTrue(content.startswith(b'%PDF-'))
                    text = ''.join(page.extract_text()
                        for page in PdfReader(BytesIO(content)).pages)
                    self.assertIn('Unit Cost', text)
                    self.assertIn('Value', text)
                    self.assertIn(f'{5:.{price_digits[1]}f}', text)
                    if by_lot:
                        self.assertIn(f'{11:.{price_digits[1]}f}', text)
                    total = 63 if by_lot else 55
                    self.assertIn(f'{total:.{price_digits[1]}f}', text)
                    if by_lot:
                        self.assertIn('LOT-VALUED', text)
                else:
                    sheet = load_workbook(BytesIO(content)).active
                    rows = list(sheet.values)
                    self.assertEqual(rows[-1][-1], 63 if by_lot else 55)
                    for row in sheet.iter_rows():
                        for cell in row[-2:]:
                            if isinstance(cell.value, (int, float)):
                                self.assertEqual(
                                    cell.number_format, value_format)
                    product_rows = [row for row in rows
                        if row[1] in (product1.rec_name, product2.rec_name)]
                    self.assertEqual(len(product_rows), 3 if by_lot else 2)
                    for row in product_rows:
                        self.assertEqual(row[-1], row[-3] * row[-2])
                    if by_lot:
                        lot_row, = [row for row in product_rows
                            if row[2] == 'LOT-VALUED']
                        self.assertEqual(lot_row[-3:], (2, 11, 22))

        # Expiry while building or saving Excel must stop the report, using
        # the same clock that started before inventory preparation.
        storage_name = storage_loc.name
        for owner, method in [(Worksheet, 'append'), (Workbook, 'save')]:
            wizard = Wizard('stock.inventory.print_total_inventory')
            wizard.form.locations.append(Location(warehouse.id))
            wizard.form.output_format = 'xlsx'
            wizard.form.timeout = 120
            original = getattr(owner, method)
            with patch.object(report_module, 'monotonic', return_value=0) as clock:
                def expire(instance, *args, **kwargs):
                    result = original(instance, *args, **kwargs)
                    if method == 'save' or (
                            args[0] and args[0][0] == storage_name):
                        clock.return_value = 120
                    return result

                with patch.object(owner, method, expire):
                    wizard.execute('print_')
            self.assertEqual(wizard.form_state, 'timed_out')
            self.assertIn('time limit exceeded', wizard.form.message)
            self.assertEqual(wizard.actions, [])
            wizard.execute('start')
            self.assertEqual(wizard.form.output_format, 'xlsx')
            self.assertEqual(wizard.form.timeout, 120)
            self.assertEqual([loc.id for loc in wizard.form.locations],
                [warehouse.id])
            wizard.execute('print_')
            self.assertEqual(wizard.state, 'end')
            self.assertEqual(len(wizard.actions), 1)
            self.assertTrue(wizard.actions[0][1].startswith(b'PK'))

        for output_format in ['pdf', 'xlsx']:
            wizard = Wizard('stock.inventory.print_total_inventory')
            wizard.form.locations.append(Location(warehouse.id))
            wizard.form.products.append(Product(product1.id))
            wizard.form.output_format = output_format
            wizard.form.valued = True
            wizard.form.group_by_lot = True
            wizard.form.order = 'product'
            wizard.form.timeout = 1
            with patch.object(report_module, 'monotonic', side_effect=[0, 1]):
                wizard.execute('print_')
            self.assertEqual(wizard.form_state, 'timed_out')
            self.assertIn('time limit exceeded', wizard.form.message)
            wizard.execute('start')
            self.assertEqual(wizard.form.output_format, output_format)
            self.assertTrue(wizard.form.valued)
            self.assertTrue(wizard.form.group_by_lot)
            self.assertEqual(wizard.form.order, 'product')
            self.assertEqual([p.id for p in wizard.form.products], [product1.id])
            wizard.form.timeout = 180
            report = (report_module.TotalInventoryXlsxReport
                if output_format == 'xlsx' else report_module.TotalInventoryReport)
            with patch.object(report, 'get_content',
                    wraps=report.get_content) as render:
                wizard.execute('print_')
            self.assertEqual(render.call_count, 1)
            self.assertEqual(wizard.state, 'end')
            self.assertEqual(len(wizard.actions), 1)
            self.assertEqual(wizard.actions[0][0], output_format)

        # Sign and product filters must still apply when products are loaded
        # from stock quantities, including an empty result.
        for quantities, expected in [('positive', 20), ('negative', 0)]:
            wizard = Wizard('stock.inventory.print_total_inventory')
            wizard.form.locations.append(Location(warehouse.id))
            wizard.form.products.append(Product(product1.id))
            wizard.form.quantities = quantities
            wizard.form.valued = True
            wizard.form.output_format = 'xlsx'
            wizard.execute('print_')
            sheet = load_workbook(BytesIO(wizard.actions[0][1])).active
            rows = list(sheet.values)
            self.assertEqual(rows[-1][-1], expected)
            if expected:
                product_row, = [r for r in rows if r[1] == product1.rec_name]
                self.assertEqual(product_row[-3:], (4, 5, 20))
                self.assertFalse(any(r[1] == product2.rec_name for r in rows))
            else:
                self.assertFalse(any(r[1] in (
                    product1.rec_name, product2.rec_name) for r in rows))
