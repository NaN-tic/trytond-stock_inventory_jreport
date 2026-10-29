
# The COPYRIGHT file at the top level of this repository contains the full
# copyright notices and license terms.
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from time import monotonic

from babel.dates import format_datetime
from fpdf import FPDF
from babel.numbers import format_decimal
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from trytond.exceptions import UserError
from trytond.model import ModelView, fields
from trytond.modules.xgettext import _

from trytond.pool import Pool, PoolMeta
from trytond.report import Report
from trytond.rpc import RPC
from trytond.tools import grouped_slice
from trytond.transaction import Transaction
from trytond.wizard import Button, StateReport, StateTransition, StateView, Wizard

class TimeoutException(UserError):
    def __init__(self):
        super().__init__(_('Inventory report time limit exceeded.'))

class TimeoutChecker:
    def __init__(self, timeout, callback):
        self._timeout = timeout
        self._callback = callback
        self._start = monotonic()

    @property
    def elapsed(self):
        return monotonic() - self._start

    def check(self):
        if self.elapsed >= self._timeout:
            raise self._callback()

class PrintTotalInventoryStart(ModelView):
    'Print Total Inventory'
    __name__ = 'stock.inventory.print_total_inventory.start'

    date = fields.Date("Date")
    products = fields.Many2Many(
        'product.product', None, None, "Products",
        domain=[
            ('type', '=', 'goods'),
        ])
    locations = fields.Many2Many(
        'stock.location', None, None, "Locations",
        domain=[('type', '=', 'warehouse')], required=True)
    output_format = fields.Selection([
        ('pdf', "PDF"),
        ('xlsx', "Excel")],
        "Format", required=True)
    valued = fields.Boolean("Valued Inventory",
        help="Show current unit costs and inventory values.")
    order = fields.Selection([
        ('location', 'Location'),
        ('product', 'Product'),
        ], "Order", required=True)
    quantities = fields.Selection([
        ('all', 'All'),
        ('positive', 'Positive'),
        ('negative', 'Negative'),
        ], "Quantities", required=True)
    timeout = fields.Integer('Timeout', required=True,
        help='Timeout in seconds')

    @staticmethod
    def default_locations():
        warehouse = Transaction().context.get('warehouse')
        if warehouse:
            return [warehouse]
        return []

    @staticmethod
    def default_output_format():
        return 'pdf'

    @staticmethod
    def default_order():
        return 'location'

    @staticmethod
    def default_quantities():
        return 'positive'

    @staticmethod
    def default_timeout():
        return 120

    @classmethod
    def __setup__(cls):
        Move = Pool().get('stock.move')
        super().__setup__()
        cls.products.domain = [
            ('type', 'in', Move.get_product_types()),
            ]
        try:
            Pool().get('stock.lot')
        except KeyError:
            return
        cls.group_by_lot = fields.Boolean('Group by Lot')

class PrintTotalInventoryTimeout(ModelView):
    'Inventory Report Timeout'
    __name__ = 'stock.inventory.print_total_inventory.timeout'

    message = fields.Text("Message", readonly=True)


class PrintTotalInventory(Wizard):
    'Print Total Inventory'
    __name__ = 'stock.inventory.print_total_inventory'

    start = StateView('stock.inventory.print_total_inventory.start',
        'stock_inventory_jreport.print_total_inventory_start_view_form', [
            Button('Cancel', 'end', 'tryton-cancel'),
            Button('Print', 'print_', 'tryton-print', default=True),
            ])
    print_ = StateTransition()
    open_ = StateReport('stock_inventory_report.total_inventory')
    timed_out = StateView('stock.inventory.print_total_inventory.timeout',
        'stock_inventory_jreport.print_total_inventory_timeout_view_form', [
            Button('Close', 'end', 'tryton-close'),
            Button('Change Parameters', 'start', 'tryton-back', default=True),
            ])

    def value_start(self, fields):
        return self.start._default_values

    def default_timed_out(self, fields):
        return {'message': _(
            'Inventory report time limit exceeded. '
            'Reduce the selection or increase the timeout and try again.')}

    def transition_print_(self):
        data = {
            'date': self.start.date,
            'valued': self.start.valued,
            'quantities': self.start.quantities,
            'group_by_lot': getattr(self.start, 'group_by_lot', False),
            'products': [x.id for x in self.start.products],
            'locations': [x.id for x in self.start.locations],
            'output_format': self.start.output_format,
            'order': self.start.order,
            'timeout': self.start.timeout,
            }
        report_name = 'stock_inventory_report.total_inventory'
        if self.start.output_format == 'xlsx':
            report_name += '_xlsx'
        Report = Pool().get(report_name, type='report')
        try:
            extension, content, direct_print, filename = Report.execute([], data)
        except TimeoutException:
            return 'timed_out'
        # Only StateView values are saved in the wizard session. Keep the
        # generated bytes on this instance until the action is returned.
        self._generated_content = content
        return 'open_'

    def do_open_(self, action):
        if self.start.output_format == 'xlsx':
            ActionReport = Pool().get('ir.action.report')
            report, = ActionReport.search([
                ('report_name', '=',
                    'stock_inventory_report.total_inventory_xlsx')])
            action = report.action.get_action_value()
        return action, {'generated_content': self._generated_content}

    def transition_open_(self):
        return 'end'

class TotalInventoryReport(Report):
    'Total Inventory Report'
    __name__ = 'stock_inventory_report.total_inventory'

    @classmethod
    def __setup__(cls):
        super().__setup__()
        cls.__rpc__['execute'] = RPC(False)

    @classmethod
    def prepare(cls, data):
        pool = Pool()
        Company = pool.get('company.company')
        Date = pool.get('ir.date')
        Location = pool.get('stock.location')
        Product = pool.get('product.product')

        checker = TimeoutChecker(data.get('timeout', 60), TimeoutException)

        if data.get('group_by_lot'):
            Lot = pool.get('stock.lot')
        else:
            Lot = None

        with Transaction().set_context(_record_cache_size=100000):
            locations = Location.search([
                    ('parent', 'child_of', data['locations']),
                    ('type', '=', 'storage'),
                    ], order=[('name', 'ASC')])
        location_ids = [l.id for l in locations]
        locations_by_id = {l.id: l for l in locations}
        checker.check()

        stock_date_end = data['date'] or Date.today()
        quantities = data.get('quantities', 'positive')

        records = []
        with Transaction().set_context(stock_date_end=stock_date_end):
            grouping = ('product',)
            if data.get('group_by_lot'):
                grouping += ('lot',)
            # Start with stock, not the entire product catalogue. In particular,
            # browse/search of a large catalogue performs expensive ORM work
            # before control can return to the timeout checker.
            pbl = Product.products_by_location(
                location_ids, grouping=grouping,
                grouping_filter=(data['products'],) if data['products'] else None)
            checker.check()
            lines = []
            product_ids = set()
            lot_ids = set()
            for key, qty in pbl.items():
                checker.check()
                if (not qty
                        or (quantities == 'positive' and qty < 0)
                        or (quantities == 'negative' and qty > 0)):
                    continue
                lines.append((key, qty))
                product_ids.add(key[1])
                if Lot and key[2]:
                    lot_ids.add(key[2])

            products_by_id = {}
            with Transaction().set_context(active_test=False):
                for ids in grouped_slice(sorted(product_ids), count=500):
                    checker.check()
                    products = Product.search([
                        ('id', 'in', list(ids)), ('type', '=', 'goods')])
                    products_by_id.update((p.id, p) for p in products)
            lots_by_id = {}
            for ids in grouped_slice(sorted(lot_ids), count=500):
                checker.check()
                lots_by_id.update((lot.id, lot) for lot in Lot.browse(ids))

            for key, qty in lines:
                checker.check()
                product = products_by_id.get(key[1])
                if product is None:
                    continue
                record = {
                    'quantity': qty,
                    'location': locations_by_id[key[0]],
                    'product': product,
                    }
                if Lot:
                    record['lot'] = lots_by_id.get(key[2])
                if data.get('valued'):
                    lot = record.get('lot')
                    cost = getattr(lot, 'cost_price', None)
                    if cost is None:
                        cost = product.cost_price
                    record['cost_price'] = cost or Decimal(0)
                    record['value'] = Decimal(str(qty)) * record['cost_price']
                records.append(record)

        company_id = Transaction().context.get('company')
        parameters = {
            'company': (Company(company_id)
                if company_id is not None and company_id >= 0 else None),
            'now': format_datetime(
                datetime.now(), format='short',
                locale=Transaction().language or 'en'),
            'sort_attribute': ('product.rec_name'
                if data['order'] == 'product' else 'location.rec_name'),
            'has_lot': data.get('group_by_lot'),
            'valued': data.get('valued', False),
            'value_format': '#,##0' + (
                '.' + '0' * Product.cost_price.digits[1]
                if Product.cost_price.digits[1] else ''),
            'timeout': data.get('timeout', 60) - checker.elapsed,
            'checker': checker,
            }
        checker.check()
        return records, parameters

    @classmethod
    def _sort_value(cls, item, sort_attribute):
        value = item
        for part in sort_attribute.split('.'):
            if isinstance(value, dict):
                value = value.get(part)
            else:
                value = getattr(value, part, '')
        return value or ''

    @classmethod
    def report_rows(cls, records, parameters):
        checker = parameters['checker']
        checker.check()
        headers = [_('Location'), _('Product')]
        if parameters.get('has_lot'):
            headers.append(_('Lot'))
        headers.append(_('Quantity'))
        if parameters.get('valued'):
            headers.extend([_('Unit Cost'), _('Value')])
        rows = []
        ordered = []
        for record in records:
            checker.check()
            ordered.append((cls._sort_value(
                record, parameters['sort_attribute']), record))
        ordered.sort(key=lambda item: item[0])
        for sort_value, record in ordered:
            checker.check()
            row = [record['location'].name, record['product'].rec_name]
            if parameters.get('has_lot'):
                lot = record.get('lot')
                row.append(lot.rec_name if lot else '')
            row.append(record['quantity'])
            if parameters.get('valued'):
                row.extend([record['cost_price'], record['value']])
            rows.append(row)
        if parameters.get('valued'):
            rows.append([_('Total')] + [''] * (len(headers) - 2) + [
                sum((r['value'] for r in records), Decimal(0))])
        checker.check()
        return headers, rows

    @classmethod
    def title(cls, action, data, records):
        parameters = data.get('parameters', {}) if data else {}
        company = parameters.get('company')
        company_name = company.rec_name if company else ''
        now = parameters.get('now', '')
        title = _('Total Inventory')
        if company_name and now:
            return '%s - %s - %s' % (title, company_name, now)
        if company_name:
            return '%s - %s' % (title, company_name)
        if now:
            return '%s - %s' % (title, now)
        return title

    @classmethod
    def get_content(cls, records, parameters):
        pdf = FPDF(orientation='L')
        fonts = Path(__file__).with_name('fonts')
        pdf.add_font('DejaVu', fname=fonts / 'DejaVuSans.ttf')
        pdf.add_font('DejaVu', style='B',
            fname=fonts / 'DejaVuSans-Bold.ttf')
        pdf.set_auto_page_break(True, margin=12)
        pdf.add_page()
        pdf.set_font('DejaVu', 'B', 12)
        pdf.multi_cell(0, 6, cls.title(None, {'parameters': parameters}, records),
            new_x='LMARGIN', new_y='NEXT')
        company = parameters.get('company')
        if parameters.get('valued') and company:
            pdf.multi_cell(0, 6, company.currency.code,
                new_x='LMARGIN', new_y='NEXT')
        pdf.ln(3)
        pdf.set_font('DejaVu', size=8)
        headers, rows = cls.report_rows(records, parameters)
        widths = [40, 85]
        if parameters.get('has_lot'):
            widths.append(40)
        widths.append(25)
        if parameters.get('valued'):
            widths.extend([30, 30])
        aligns = ['LEFT'] * (len(headers) - (3 if parameters.get('valued') else 1))
        aligns += ['RIGHT'] * (3 if parameters.get('valued') else 1)
        checker = parameters['checker']
        checker.check()
        with pdf.table(col_widths=widths, text_align=aligns,
                line_height=5) as table:
            table.row(headers)
            for values in rows:
                checker.check()
                table.row([format_decimal(value,
                    format=(parameters['value_format']
                        if parameters.get('valued')
                        and index >= len(values) - 2 else '#,##0.######'),
                    locale=Transaction().language or 'en')
                    if isinstance(value, (Decimal, float, int)) else value
                    for index, value in enumerate(values)])
        if not records:
            pdf.multi_cell(0, 6, _('No records found'))
        checker.check()
        content = bytes(pdf.output())
        checker.check()
        return content

    @classmethod
    def execute(cls, ids, data):
        ActionReport = Pool().get('ir.action.report')
        action, = ActionReport.search([('report_name', '=', cls.__name__)])
        cls.check_access(action, action.model, ids)
        if 'generated_content' in data:
            return ('pdf', bytes(data['generated_content']),
                action.direct_print, action.name)
        with Transaction().set_context(active_test=False):
            records, parameters = cls.prepare(data)
            content = cls.get_content(records, parameters)
        return 'pdf', content, action.direct_print, action.name

class TotalInventoryXlsxReport(Report, metaclass=PoolMeta):
    __name__ = 'stock_inventory_report.total_inventory_xlsx'

    @classmethod
    def execute(cls, ids, data):
        pool = Pool()
        ActionReport = pool.get('ir.action.report')
        action_report, = ActionReport.search([
                ('report_name', '=', cls.__name__)
                ])
        cls.check_access(action_report, action_report.model, ids)
        if 'generated_content' in data:
            return ('xlsx', bytes(data['generated_content']),
                action_report.direct_print, action_report.name)
        with Transaction().set_context(active_test=False):
            records, parameters = TotalInventoryReport.prepare(data)
        content = cls.get_content(records, parameters)
        filename = action_report.name
        return 'xlsx', content, action_report.direct_print, filename

    @classmethod
    def get_content(cls, records, parameters):
        checker = parameters['checker']
        checker.check()
        wb = Workbook()
        ws = wb.active
        ws.title = _('Total Inventory')[:31]

        title = TotalInventoryReport.title(
            None, {'parameters': parameters, 'records': records}, records)
        ws.append([title])
        company = parameters.get('company')
        if company:
            ws.append([company.rec_name])
        if parameters.get('now'):
            ws.append([parameters['now']])
        ws.append([])

        headers, rows = TotalInventoryReport.report_rows(records, parameters)
        company = parameters.get('company')
        if parameters.get('valued') and company:
            ws.append([company.currency.code])
        ws.append(headers)
        header_row = ws.max_row
        for cell in ws[header_row]:
            cell.font = Font(bold=True)
        for row_number, row in enumerate(rows, header_row + 1):
            checker.check()
            ws.append(row)
            for cell in ws[row_number]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                elif isinstance(cell.value, (Decimal, float, int)):
                    cell.number_format = (parameters['value_format']
                        if parameters.get('valued')
                        and cell.column >= len(headers) - 1 else '#,##0.######')
        if parameters.get('valued'):
            for cell in ws[header_row + len(rows)]:
                cell.font = Font(bold=True)
        for index, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(index)].width = (
                55 if index == 2 else 22)
        ws.freeze_panes = f'A{header_row + 1}'
        ws.auto_filter.ref = f'A{header_row}:{get_column_letter(len(headers))}' \
            f'{header_row + len(records)}'
        stream = BytesIO()
        checker.check()
        wb.save(stream)
        checker.check()
        return stream.getvalue()
